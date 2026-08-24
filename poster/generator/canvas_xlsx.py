# -*- coding: utf-8 -*-
"""Мини-«холст» поверх листа Excel.

Лист превращается в сетку квадратных ячеек по 5 мм: столбцы получают
одинаковую ширину, строки — одинаковую высоту. Дальше по этой сетке можно
рисовать заливками (стены, стрелки, секторы) и подписывать текстом
в объединённых диапазонах.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Ячейка 5 мм: 19 пикселей по ширине (≈5,03 мм) и 14,25 пункта по высоте.
CELL_MM = 19 * 25.4 / 96.0
COL_WIDTH = (19 - 5) / 7.0
ROW_HEIGHT = 19 * 72.0 / 96.0

FONT_NAME = "Arial"


class XlsxCanvas:
    def __init__(self, ws, cols, rows, background="FFFFFF"):
        self.ws = ws
        self.cols = cols
        self.rows = rows
        self._fills = {}
        self._fonts = {}
        self._merged = set()

        ws.sheet_view.showGridLines = False
        for c in range(1, cols + 1):
            ws.column_dimensions[get_column_letter(c)].width = COL_WIDTH
        for r in range(1, rows + 1):
            ws.row_dimensions[r].height = ROW_HEIGHT

        if background:
            self.fill(0, 0, cols - 1, rows - 1, background)

    # -- служебное ---------------------------------------------------------

    def _fill_obj(self, color):
        if color not in self._fills:
            self._fills[color] = PatternFill("solid", fgColor=color)
        return self._fills[color]

    def _font_obj(self, size, bold, color, italic=False, name=FONT_NAME):
        key = (size, bold, color, italic, name)
        if key not in self._fonts:
            self._fonts[key] = Font(name=name, size=size, bold=bold,
                                    color=color, italic=italic)
        return self._fonts[key]

    def cell(self, col, row):
        return self.ws.cell(row=row + 1, column=col + 1)

    def _inside(self, col, row):
        return 0 <= col < self.cols and 0 <= row < self.rows

    # -- заливки -----------------------------------------------------------

    def dot(self, col, row, color):
        if self._inside(col, row):
            self.cell(col, row).fill = self._fill_obj(color)

    def fill(self, c1, r1, c2, r2, color):
        c1, c2 = min(c1, c2), max(c1, c2)
        r1, r2 = min(r1, r2), max(r1, r2)
        f = self._fill_obj(color)
        for r in range(max(r1, 0), min(r2, self.rows - 1) + 1):
            for c in range(max(c1, 0), min(c2, self.cols - 1) + 1):
                self.cell(c, r).fill = f

    def polygon(self, points, color):
        """Заливка произвольного многоугольника (координаты — в ячейках)."""
        if len(points) < 3:
            return
        ys = [p[1] for p in points]
        r_min = max(int(min(ys)), 0)
        r_max = min(int(max(ys)) + 1, self.rows - 1)
        n = len(points)
        for r in range(r_min, r_max + 1):
            yc = r + 0.5
            xs = []
            for i in range(n):
                x1, y1 = points[i]
                x2, y2 = points[(i + 1) % n]
                if (y1 <= yc < y2) or (y2 <= yc < y1):
                    t = (yc - y1) / (y2 - y1)
                    xs.append(x1 + t * (x2 - x1))
            xs.sort()
            for i in range(0, len(xs) - 1, 2):
                c1 = int(round(xs[i]))
                c2 = int(round(xs[i + 1])) - 1
                if c2 >= c1:
                    self.fill(c1, r, c2, r, color)

    # -- линии и стрелки ---------------------------------------------------

    @staticmethod
    def _segment_cells(p1, p2):
        (c1, r1), (c2, r2) = p1, p2
        if c1 == c2:
            step = 1 if r2 >= r1 else -1
            return [(c1, r) for r in range(r1, r2 + step, step)]
        if r1 == r2:
            step = 1 if c2 >= c1 else -1
            return [(c, r1) for c in range(c1, c2 + step, step)]
        # диагональ — растеризуем по Брезенхэму
        cells = []
        dc, dr = c2 - c1, r2 - r1
        n = max(abs(dc), abs(dr))
        for i in range(n + 1):
            cells.append((c1 + round(dc * i / n), r1 + round(dr * i / n)))
        return cells

    def polyline(self, points, color, width=1, dash=None,
                 head=0, head_size=2):
        """Ортогональная (или диагональная) ломаная заданной толщины."""
        off = (width - 1) // 2
        counter = 0
        for p1, p2 in zip(points, points[1:]):
            horizontal = p1[1] == p2[1]
            for (c, r) in self._segment_cells(p1, p2):
                draw = True
                if dash:
                    on, gap = dash
                    draw = (counter % (on + gap)) < on
                counter += 1
                if not draw:
                    continue
                if horizontal:
                    self.fill(c, r - off, c, r - off + width - 1, color)
                else:
                    self.fill(c - off, r, c - off + width - 1, r, color)
        if head:
            self.arrow_head(points[-2], points[-1], color, head_size)

    def arrow_head(self, p_from, p_to, color, size=2):
        """Треугольный наконечник: вершина в p_to, основание шириной 2*size+1."""
        c, r = p_to
        dc, dr = p_to[0] - p_from[0], p_to[1] - p_from[1]
        for i in range(size + 1):
            if abs(dc) >= abs(dr):
                cc = c - i if dc > 0 else c + i
                self.fill(cc, r - i, cc, r + i, color)
            else:
                rr = r - i if dr > 0 else r + i
                self.fill(c - i, rr, c + i, rr, color)

    # -- рамки -------------------------------------------------------------

    def outline(self, c1, r1, c2, r2, color="000000", style="thin"):
        side = Side(style=style, color=color)
        for c in range(c1, c2 + 1):
            top = self.cell(c, r1)
            top.border = self._merge_border(top.border, top=side)
            bot = self.cell(c, r2)
            bot.border = self._merge_border(bot.border, bottom=side)
        for r in range(r1, r2 + 1):
            left = self.cell(c1, r)
            left.border = self._merge_border(left.border, left=side)
            right = self.cell(c2, r)
            right.border = self._merge_border(right.border, right=side)

    @staticmethod
    def _merge_border(existing, top=None, bottom=None, left=None, right=None):
        return Border(
            top=top or existing.top,
            bottom=bottom or existing.bottom,
            left=left or existing.left,
            right=right or existing.right,
        )

    # -- текст -------------------------------------------------------------

    def text(self, c1, r1, c2, r2, value, size=11, bold=False,
             color="000000", fill=None, align="left", valign="center",
             wrap=False, rotation=0, italic=False, indent=0, name=FONT_NAME):
        if (c1, r1, c2, r2) not in self._merged and (c1, r1) != (c2, r2):
            self.ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1,
                                end_row=r2 + 1, end_column=c2 + 1)
            self._merged.add((c1, r1, c2, r2))
        cell = self.cell(c1, r1)
        cell.value = value
        cell.font = self._font_obj(size, bold, color, italic, name)
        cell.alignment = Alignment(horizontal=align, vertical=valign,
                                   wrap_text=wrap, text_rotation=rotation,
                                   indent=indent)
        if fill:
            self.fill(c1, r1, c2, r2, fill)
        return cell

    def lines(self, c1, r1, c2, texts, size=10, color="1B2430",
              bold=False, step=1, align="left", indent=0):
        """Столбик строк текста, по одной строке на ряд ячеек."""
        for i, t in enumerate(texts):
            r = r1 + i * step
            self.text(c1, r, c2, r + step - 1, t, size=size, bold=bold,
                      color=color, align=align, indent=indent)
        return r1 + len(texts) * step

    def panel(self, c1, r1, c2, r2, title=None, title_bg="1B2430",
              title_fg="FFFFFF", body_bg="FFFFFF", border="1B2430",
              title_size=12, title_rows=3):
        self.fill(c1, r1, c2, r2, body_bg)
        if title is not None:
            self.fill(c1, r1, c2, r1 + title_rows - 1, title_bg)
            self.text(c1, r1, c2, r1 + title_rows - 1, "  " + title,
                      size=title_size, bold=True, color=title_fg,
                      align="left", valign="center")
        self.outline(c1, r1, c2, r2, color=border, style="medium")
        return r1 + (title_rows if title is not None else 0)
